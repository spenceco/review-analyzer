"""Build the downloadable, formatted Excel report with openpyxl."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .models import SENTIMENTS, AnalysisResult, Insight

# Same palette as the web page: blue = positive, grey = neutral, red = negative.
SENTIMENT_COLORS = {"positive": "2A78D6", "neutral": "A3A29B", "negative": "E34948"}
SENTIMENT_TINTS = {"positive": "DCEAFB", "neutral": "F0EFEC", "negative": "FBDCDB"}
INK = "0B0B0B"
MUTED = "6B6A65"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HAIRLINE = Border(bottom=Side(style="thin", color="E1E0D9"))

TITLE = Font(name="Calibri", size=18, bold=True, color=INK)
H2 = Font(name="Calibri", size=13, bold=True, color=INK)
BOLD = Font(name="Calibri", bold=True, color=INK)
SMALL_MUTED = Font(name="Calibri", size=10, color=MUTED)
WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def _header_row(ws: Worksheet, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = WHITE_BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")


def _insight_section(ws: Worksheet, row: int, heading: str, items: list[Insight]) -> int:
    ws.cell(row=row, column=1, value=heading).font = H2
    row += 1
    _header_row(ws, row, ["#", "Theme", "Details", "Reviews"])
    row += 1
    if not items:
        ws.cell(row=row, column=2, value="None found").font = SMALL_MUTED
        return row + 2
    for i, item in enumerate(items, start=1):
        values = [i, item.title, item.detail, item.mentions]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = WRAP_TOP
            cell.border = HAIRLINE
        ws.cell(row=row, column=2).font = BOLD
        row += 1
    return row + 1


def _summary_sheet(ws: Worksheet, result: AnalysisResult) -> None:
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    for letter, width in zip("ABCDEF", (6, 30, 70, 10, 4, 14), strict=True):
        ws.column_dimensions[letter].width = width

    ws["A1"] = "Customer review analysis"
    ws["A1"].font = TITLE
    generated = datetime.now(UTC).strftime("%d %b %Y %H:%M UTC")
    ws["A2"] = f"{result.filename} · {len(result.reviews)} reviews · generated {generated}"
    ws["A2"].font = SMALL_MUTED

    # Sentiment table (also the data source for the doughnut chart).
    total = max(1, len(result.analyses))
    ws["A4"] = "Sentiment"
    ws["A4"].font = H2
    _header_row(ws, 5, ["", "Sentiment", "Share", "Reviews"])
    for i, s in enumerate(SENTIMENTS):
        r = 6 + i
        count = result.sentiment_counts.get(s, 0)
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=SENTIMENT_COLORS[s])
        ws.cell(row=r, column=2, value=s.capitalize()).font = BOLD
        share = ws.cell(row=r, column=3, value=count / total)
        share.number_format = "0%"
        share.alignment = Alignment(horizontal="left")
        ws.cell(row=r, column=4, value=count)
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = HAIRLINE
    ratings = [r.rating for r in result.reviews if r.rating is not None]
    if ratings:
        ws["B10"] = "Average rating"
        ws["B10"].font = BOLD
        ws["C10"] = round(sum(ratings) / len(ratings), 2)
        ws["C10"].alignment = Alignment(horizontal="left")

    chart = DoughnutChart(holeSize=60)
    chart.title = None
    chart.add_data(Reference(ws, min_col=4, min_row=6, max_row=8), titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=2, min_row=6, max_row=8))
    series = chart.series[0]
    for idx, s in enumerate(SENTIMENTS):
        point = DataPoint(idx=idx)
        point.graphicalProperties.solidFill = SENTIMENT_COLORS[s]
        point.graphicalProperties.line.solidFill = "FFFFFF"
        series.dPt.append(point)
    chart.dataLabels = DataLabelList(showPercent=True, showVal=False, showCatName=False,
                                     showSerName=False, showLegendKey=False)
    chart.legend.position = "r"
    chart.height, chart.width = 6.5, 9.5
    ws.add_chart(chart, "F4")

    row = 13
    ws.cell(row=row, column=1, value="Suggested actions").font = H2
    row += 1
    _header_row(ws, row, ["#", "Action", "Why / first step"])
    row += 1
    for i, action in enumerate(result.insights.suggested_actions, start=1):
        for col, value in enumerate([i, action.title, action.detail], start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = WRAP_TOP
            cell.border = HAIRLINE
        ws.cell(row=row, column=2).font = BOLD
        row += 1
    row += 1
    row = _insight_section(ws, row, "Top complaints", result.insights.top_complaints)
    _insight_section(ws, row, "Top praise", result.insights.top_praise)


def _topics_sheet(ws: Worksheet, result: AnalysisResult) -> None:
    ws.sheet_view.showGridLines = False
    headers = ["Topic", "Positive", "Neutral", "Negative", "Total"]
    _header_row(ws, 1, headers)
    for r, stat in enumerate(result.topic_stats, start=2):
        for c, value in enumerate(
            [stat.topic, stat.positive, stat.neutral, stat.negative, stat.total], start=1
        ):
            ws.cell(row=r, column=c, value=value).border = HAIRLINE
    ws.column_dimensions["A"].width = 26
    for col in "BCDE":
        ws.column_dimensions[col].width = 11
    last = len(result.topic_stats) + 1
    if last < 2:
        return

    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.gapWidth = 60
    chart.title = "Reviews per topic, by sentiment"
    chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=1, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    for series, s in zip(chart.series, SENTIMENTS, strict=True):
        series.graphicalProperties.solidFill = SENTIMENT_COLORS[s]
        series.graphicalProperties.line.solidFill = "FFFFFF"
    chart.y_axis.majorGridlines = None
    chart.x_axis.scaling.orientation = "maxMin"  # Biggest topic at the top.
    chart.legend.position = "b"
    chart.height, chart.width = 9, 16
    ws.add_chart(chart, "G2")


def _reviews_sheet(ws: Worksheet, result: AnalysisResult) -> None:
    headers = ["#", "Date", "Rating", "Sentiment", "Topic", "Summary", "Review"]
    ws.append(headers)
    by_id = {a.id: a for a in result.analyses}
    for review in result.reviews:
        a = by_id[review.id]
        ws.append([review.id, review.date, review.rating, a.sentiment, a.topic, a.summary,
                   review.text])
    for col, width in zip(range(1, 8), (6, 12, 8, 12, 20, 50, 70), strict=True):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
        sentiment = row[3].value
        if isinstance(sentiment, str) and sentiment in SENTIMENT_TINTS:
            row[3].fill = PatternFill("solid", fgColor=SENTIMENT_TINTS[sentiment])
            color = MUTED if sentiment == "neutral" else SENTIMENT_COLORS[sentiment]
            row[3].font = Font(name="Calibri", bold=True, color=color)
    ref = f"A1:G{len(result.reviews) + 1}"
    table = Table(displayName="Reviews", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)  # Adds sortable/filterable headers.
    _header_row(ws, 1, headers)
    ws.freeze_panes = "A2"


def _print_setup(ws: Worksheet) -> None:
    """Landscape, one page wide, so each sheet prints cleanly."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4


def build_report(result: AnalysisResult) -> bytes:
    """Return the .xlsx file as bytes."""
    wb = Workbook()
    _summary_sheet(wb.active, result)
    _topics_sheet(wb.create_sheet("Topics"), result)
    _reviews_sheet(wb.create_sheet("Reviews"), result)
    for ws in wb.worksheets:
        _print_setup(ws)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
