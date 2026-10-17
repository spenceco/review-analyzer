"""End-to-end tests through the HTTP API with the offline provider."""

from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings
from app.main import create_app

from .conftest import wait_for_job


def upload(client: TestClient, data: bytes, name: str = "reviews.csv"):  # type: ignore[no-untyped-def]
    return client.post("/api/analyze", files={"file": (name, data, "text/csv")})


def test_index_and_sample_served(client: TestClient) -> None:
    page = client.get("/")
    assert page.status_code == 200 and "Review Analyzer" in page.text
    assert "default-src 'self'" in page.headers["content-security-policy"]
    assert client.get("/api/sample.csv").text.startswith("date,rating,review")


def test_status_reports_limits(client: TestClient) -> None:
    status = client.get("/api/status").json()
    assert status["can_analyze"] is True
    assert status["max_rows"] == 200 and status["calls_remaining_today"] == 50
    assert status["uploads_remaining_for_you"] == 3


def test_full_flow_with_excel_report(client: TestClient, sample_bytes: bytes) -> None:
    started = upload(client, sample_bytes, "coffee.csv")
    assert started.status_code == 202, started.text
    body = started.json()
    assert body["total_reviews"] == 75 and body["estimated_calls"] == 4

    job = wait_for_job(client, body["job_id"])
    assert job["status"] == "done", job
    result = job["result"]
    assert len(result["analyses"]) == 75
    assert sum(result["sentiment_counts"].values()) == 75
    assert len(result["insights"]["suggested_actions"]) == 3
    assert result["insights"]["top_complaints"][0]["mentions"] > 0
    assert client.get("/api/status").json()["calls_remaining_today"] == 46

    report = client.get(f"/api/jobs/{body['job_id']}/report.xlsx")
    assert report.status_code == 200
    assert 'filename="coffee-analysis.xlsx"' in report.headers["content-disposition"]
    wb = load_workbook(BytesIO(report.content))
    assert wb.sheetnames == ["Summary", "Topics", "Reviews"]
    assert wb["Reviews"].max_row == 76


def test_invalid_csv_returns_helpful_message(client: TestClient) -> None:
    res = upload(client, b"name,email\nBob,bob@example.com\n")
    assert res.status_code == 422
    assert "Couldn't find a review column" in res.json()["detail"]["message"]


def test_rejects_non_csv_extension(client: TestClient) -> None:
    res = upload(client, b"review\nhello there\n", "reviews.xlsx")
    assert res.status_code == 400 and res.json()["detail"]["code"] == "bad_type"


def test_oversized_upload_rejected(client: TestClient) -> None:
    res = upload(client, b"review\n" + b"x" * (3 * 1024 * 1024))
    assert res.status_code in (413, 422)


def test_rate_limit_per_ip(client: TestClient) -> None:
    for _ in range(3):
        job_id = upload(client, b"review\nGreat coffee\n").json()["job_id"]
        wait_for_job(client, job_id)
    res = upload(client, b"review\nGreat coffee\n")
    assert res.status_code == 429 and res.json()["detail"]["code"] == "rate_limited"
    assert client.get("/api/status").json()["uploads_remaining_for_you"] == 0


def test_invalid_uploads_do_not_use_rate_limit(client: TestClient) -> None:
    for _ in range(5):
        assert upload(client, b"nope\n").status_code == 422
    assert client.get("/api/status").json()["uploads_remaining_for_you"] == 3


def test_daily_limit_disables_analysis(settings: Settings, sample_bytes: bytes) -> None:
    settings.daily_api_call_limit = 3
    with TestClient(create_app(settings)) as client:
        # 75 reviews need 4 calls but only 3 are left: refused before spending anything.
        res = upload(client, sample_bytes)
        assert res.status_code == 429 and "smaller file" in res.json()["detail"]["message"]

        wait_for_job(client, upload(client, b"review\nGreat coffee\n").json()["job_id"])
        status = client.get("/api/status").json()
        assert status["calls_remaining_today"] == 1
        assert status["can_analyze"] is False and "daily" in status["reason"]
        res = upload(client, b"review\nGreat coffee\n")
        assert res.status_code == 429 and res.json()["detail"]["code"] == "daily_limit"


def test_missing_api_key_explains_itself(settings: Settings) -> None:
    settings.ai_provider = "gemini"
    settings.gemini_api_key = None
    with TestClient(create_app(settings)) as client:
        status = client.get("/api/status").json()
        assert status["can_analyze"] is False and "configured" in status["reason"]
        assert upload(client, b"review\nhello there\n").status_code == 503


def test_unknown_job(client: TestClient) -> None:
    assert client.get("/api/jobs/nope").status_code == 404
    assert client.get("/api/jobs/nope/report.xlsx").status_code == 404
