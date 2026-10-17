from __future__ import annotations

import asyncio
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.analyzer import Analyzer
from app.csv_validation import parse_reviews
from app.limits import DailyBudget
from app.providers.fake import FakeProvider
from app.report import build_report


def test_report_contents(tmp_path: Path, sample_bytes: bytes) -> None:
    reviews = parse_reviews(sample_bytes, max_rows=200, max_bytes=2_000_000,
                            max_review_chars=1500)
    analyzer = Analyzer(FakeProvider(), DailyBudget(tmp_path / "u.db", 100))
    result = asyncio.run(analyzer.analyze(reviews, filename="coffee.csv"))
    wb = load_workbook(BytesIO(build_report(result)))

    summary = wb["Summary"]
    assert summary["A1"].value == "Customer review analysis"
    assert [summary.cell(row=r, column=2).value for r in (6, 7, 8)] == \
        ["Positive", "Neutral", "Negative"]
    assert sum(summary.cell(row=r, column=4).value for r in (6, 7, 8)) == 75
    assert summary["B10"].value == "Average rating"
    assert len(summary._charts) == 1

    topics = wb["Topics"]
    assert topics["A1"].value == "Topic" and len(topics._charts) == 1

    rows = wb["Reviews"]
    assert rows.freeze_panes == "A2" and "Reviews" in rows.tables
    assert rows.cell(row=2, column=7).value == reviews[0].text
